"""
importador.py — Lectura de Excel / CSV y mapeo automático de columnas.

No depende de Streamlit ni de Supabase: recibe un archivo, devuelve un
DataFrame limpio y una lista de avisos. Así es fácil de probar.

Qué resuelve:
  * Encuentra la fila de encabezados aunque el archivo tenga un título,
    un logo o filas vacías arriba (típico de las exportaciones de un POS).
  * Reconoce los nombres de columna más habituales en español e inglés,
    sin importar mayúsculas, tildes ni espacios.
  * Interpreta números en formato local: "$ 1.234,50", "1,234.50", "12,5".
  * Avisa de filas incompletas, duplicados y precios menores al costo,
    en lugar de romperse o de guardar basura en silencio.
"""

from __future__ import annotations

import io
import re
import unicodedata

import pandas as pd

CAMPOS_OBLIGATORIOS = ("item_name", "cost", "price", "units_sold")
CAMPOS_OPCIONALES = ("category",)
TODOS_LOS_CAMPOS = CAMPOS_OBLIGATORIOS + CAMPOS_OPCIONALES

ETIQUETAS = {
    "item_name": "Plato",
    "category": "Categoría",
    "cost": "Costo",
    "price": "Precio",
    "units_sold": "Unidades vendidas",
}

#: Nombres de columna que se reconocen automáticamente para cada campo.
SINONIMOS: dict[str, tuple[str, ...]] = {
    "item_name": (
        "plato", "platos", "nombre", "nombre del plato", "nombre plato",
        "producto", "productos", "item", "items", "articulo", "artículo",
        "descripcion", "descripción", "detalle", "menu item", "item name",
        "platillo", "comida", "receta",
    ),
    "category": (
        "categoria", "categoría", "rubro", "familia", "grupo", "tipo",
        "seccion", "sección", "category", "group", "linea", "línea",
    ),
    "cost": (
        "costo", "costos", "costo unitario", "costo plato", "costo por plato",
        "food cost", "cmv", "cost", "unit cost", "costo materia prima",
        "costo directo",
    ),
    "price": (
        "precio", "precios", "precio de venta", "precio venta", "pvp",
        "precio unitario", "price", "selling price", "unit price", "venta",
        "precio carta",
    ),
    "units_sold": (
        "unidades", "unidades vendidas", "cantidad", "cantidades", "cant",
        "ventas", "vendidos", "units", "units sold", "qty", "quantity",
        "cubiertos", "porciones", "volumen",
    ),
}


# ---------------------------------------------------------------------------
# Normalización de textos
# ---------------------------------------------------------------------------
def _norm(texto) -> str:
    """minúsculas, sin tildes, sin signos: 'Precio de Venta ($)' -> 'precio de venta'."""
    if texto is None:
        return ""
    s = unicodedata.normalize("NFKD", str(texto).strip().lower())
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


_SINONIMOS_NORM = {
    campo: {_norm(x) for x in nombres} for campo, nombres in SINONIMOS.items()
}
_TODOS_SINONIMOS = set().union(*_SINONIMOS_NORM.values())


# ---------------------------------------------------------------------------
# Números en formato local
# ---------------------------------------------------------------------------
def parsear_numero(valor) -> float | None:
    """Convierte a float textos como '$ 1.234,50', '1,234.50', '12,5', '  7 '.

    Criterio: si aparecen los dos separadores, el ÚLTIMO es el decimal.
    Si aparece solo la coma, es decimal (uso habitual en el Río de la Plata).
    Si aparece solo el punto una vez, es decimal; si aparece varias veces,
    son separadores de miles.
    """
    if valor is None:
        return None
    if isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return None if pd.isna(valor) else float(valor)

    s = str(valor).strip()
    if not s or _norm(s) in ("", "nan", "none", "na", "s d", "sd"):
        return None

    negativo = s.startswith("-") or (s.startswith("(") and s.endswith(")"))
    s = re.sub(r"[^\d,.]", "", s)
    if not s or not any(c.isdigit() for c in s):
        return None

    tiene_coma, tiene_punto = "," in s, "." in s
    if tiene_coma and tiene_punto:
        if s.rfind(",") > s.rfind("."):          # 1.234,50
            s = s.replace(".", "").replace(",", ".")
        else:                                     # 1,234.50
            s = s.replace(",", "")
    elif tiene_coma:
        s = s.replace(",", ".") if s.count(",") == 1 else s.replace(",", "")
    elif s.count(".") > 1:                        # 1.234.567
        s = s.replace(".", "")

    try:
        numero = float(s)
    except ValueError:
        return None
    return -numero if negativo else numero


# ---------------------------------------------------------------------------
# Lectura del archivo
# ---------------------------------------------------------------------------
def listar_hojas(archivo) -> list[str]:
    """Nombres de las hojas de un Excel. Lista vacía para un CSV."""
    nombre = getattr(archivo, "name", "").lower()
    if nombre.endswith(".csv") or nombre.endswith(".txt"):
        return []
    archivo.seek(0)
    with pd.ExcelFile(archivo) as libro:
        return list(libro.sheet_names)


def _detectar_separador(texto: str) -> str:
    primera = texto.splitlines()[0] if texto.splitlines() else ""
    conteos = {sep: primera.count(sep) for sep in (";", ",", "\t", "|")}
    mejor = max(conteos, key=conteos.get)
    return mejor if conteos[mejor] > 0 else ","


def _fila_encabezado(crudo: pd.DataFrame) -> int:
    """Busca en las primeras filas cuál parece ser la de los encabezados."""
    limite = min(15, len(crudo))
    mejor_fila, mejor_puntaje = 0, 0
    for i in range(limite):
        celdas = [_norm(v) for v in crudo.iloc[i].tolist()]
        puntaje = sum(1 for c in celdas if c and c in _TODOS_SINONIMOS)
        if puntaje > mejor_puntaje:
            mejor_fila, mejor_puntaje = i, puntaje
    return mejor_fila if mejor_puntaje >= 2 else 0


def leer(archivo, hoja: str | None = None) -> pd.DataFrame:
    """Devuelve el contenido del archivo con los encabezados ya ubicados."""
    nombre = getattr(archivo, "name", "").lower()
    archivo.seek(0)

    if nombre.endswith((".csv", ".txt")):
        datos = archivo.read()
        if isinstance(datos, str):
            texto = datos
        else:
            for codificacion in ("utf-8-sig", "utf-8", "latin-1"):
                try:
                    texto = datos.decode(codificacion)
                    break
                except UnicodeDecodeError:
                    continue
            else:
                raise ValueError("No pude leer el archivo: codificación desconocida.")
        crudo = pd.read_csv(
            io.StringIO(texto), sep=_detectar_separador(texto),
            header=None, dtype=object, engine="python",
        )
    else:
        crudo = pd.read_excel(
            archivo, sheet_name=hoja if hoja else 0, header=None, dtype=object
        )

    if crudo.empty:
        return pd.DataFrame()

    fila = _fila_encabezado(crudo)
    encabezados = [
        str(v).strip() if v is not None and not pd.isna(v) else f"columna_{i + 1}"
        for i, v in enumerate(crudo.iloc[fila].tolist())
    ]

    df = crudo.iloc[fila + 1:].copy()
    df.columns = _unicos(encabezados)
    df = df.dropna(how="all").reset_index(drop=True)
    df = df.loc[:, [c for c in df.columns if not str(c).startswith("columna_")
                    or df[c].notna().any()]]
    return df


def _unicos(nombres: list[str]) -> list[str]:
    vistos: dict[str, int] = {}
    salida = []
    for n in nombres:
        if n in vistos:
            vistos[n] += 1
            salida.append(f"{n}_{vistos[n]}")
        else:
            vistos[n] = 0
            salida.append(n)
    return salida


# ---------------------------------------------------------------------------
# Mapeo de columnas
# ---------------------------------------------------------------------------
def detectar_columnas(df: pd.DataFrame) -> dict[str, str | None]:
    """Adivina qué columna del archivo corresponde a cada campo."""
    normales = {col: _norm(col) for col in df.columns}
    usados: set = set()
    mapeo: dict[str, str | None] = {}

    for campo in TODOS_LOS_CAMPOS:
        sinonimos = _SINONIMOS_NORM[campo]
        elegido = None

        for col, n in normales.items():          # coincidencia exacta
            if col not in usados and n in sinonimos:
                elegido = col
                break

        if elegido is None:                      # coincidencia parcial
            for col, n in normales.items():
                if col in usados or not n:
                    continue
                if any(s in n or n in s for s in sinonimos if len(s) >= 4):
                    elegido = col
                    break

        mapeo[campo] = elegido
        if elegido:
            usados.add(elegido)
    return mapeo


# ---------------------------------------------------------------------------
# Preparación y validación
# ---------------------------------------------------------------------------
def preparar(
    df: pd.DataFrame, mapeo: dict[str, str | None]
) -> tuple[pd.DataFrame, list[tuple[str, str]]]:
    """Arma el DataFrame final y devuelve los mensajes para mostrar.

    Cada mensaje es ('error' | 'aviso', texto). Con un 'error' no se puede
    continuar; un 'aviso' es informativo y no impide guardar.
    """
    mensajes: list[tuple[str, str]] = []

    faltan = [ETIQUETAS[c] for c in CAMPOS_OBLIGATORIOS if not mapeo.get(c)]
    if faltan:
        mensajes.append(
            ("error", "Falta indicar qué columna corresponde a: " + ", ".join(faltan) + ".")
        )
        return pd.DataFrame(), mensajes
    if df.empty:
        mensajes.append(("error", "El archivo no tiene filas de datos."))
        return pd.DataFrame(), mensajes

    out = pd.DataFrame(index=df.index)
    out["item_name"] = df[mapeo["item_name"]].map(
        lambda v: "" if v is None or pd.isna(v) else str(v).strip()
    )
    if mapeo.get("category"):
        out["category"] = df[mapeo["category"]].map(
            lambda v: "" if v is None or pd.isna(v) else str(v).strip()
        )
    else:
        out["category"] = ""
        mensajes.append((
            "aviso",
            "El archivo no trae columna de categoría: todos los platos van a "
            "quedar en «Sin categoría», y la matriz se calcula comparando entre sí "
            "platos que quizá no sean comparables.",
        ))

    for campo in ("cost", "price", "units_sold"):
        out[campo] = df[mapeo[campo]].map(parsear_numero)

    # Filas sin nombre de plato: son separadores, subtotales o filas vacías.
    vacias = out["item_name"].map(lambda s: s == "" or _norm(s) in ("nan", "total", "totales"))
    descartadas = int(vacias.sum())
    out = out[~vacias]
    if descartadas:
        mensajes.append(
            ("aviso", f"Descarté {descartadas} fila(s) sin nombre de plato (vacías, totales o separadores).")
        )

    if out.empty:
        mensajes.append(("error", "Después de limpiar, no quedó ninguna fila válida."))
        return pd.DataFrame(), mensajes

    # Filas con números ilegibles.
    incompletas = out[["cost", "price", "units_sold"]].isna().any(axis=1)
    if incompletas.any():
        nombres = out.loc[incompletas, "item_name"].head(5).tolist()
        extra = "…" if int(incompletas.sum()) > 5 else ""
        mensajes.append((
            "aviso",
            f"Salteé {int(incompletas.sum())} fila(s) porque no pude leer el costo, "
            f"el precio o las unidades: {', '.join(nombres)}{extra}",
        ))
        out = out[~incompletas]

    if out.empty:
        mensajes.append(("error", "Ninguna fila tenía costo, precio y unidades legibles."))
        return pd.DataFrame(), mensajes

    # Negativos -> 0
    for campo in ("cost", "price", "units_sold"):
        negativos = out[campo] < 0
        if negativos.any():
            mensajes.append(
                ("aviso", f"{int(negativos.sum())} valor(es) negativo(s) en {ETIQUETAS[campo]}: los puse en 0.")
            )
            out.loc[negativos, campo] = 0

    out["units_sold"] = out["units_sold"].round().astype(int)
    out["cost"] = out["cost"].round(2)
    out["price"] = out["price"].round(2)
    out["category"] = out["category"].replace("", "Sin categoría")

    # Duplicados: el upsert falla si el mismo plato viene dos veces en el lote.
    clave = out["item_name"].map(_norm)
    duplicados = clave.duplicated(keep="last")
    if duplicados.any():
        repetidos = out.loc[duplicados, "item_name"].head(5).tolist()
        mensajes.append((
            "aviso",
            f"{int(duplicados.sum())} plato(s) repetido(s) ({', '.join(repetidos)}): "
            "me quedo con la última aparición de cada uno.",
        ))
        out = out[~duplicados]

    # Señales de que el archivo puede estar mal mapeado.
    a_perdida = out["price"] < out["cost"]
    if a_perdida.any():
        nombres = out.loc[a_perdida, "item_name"].head(5).tolist()
        mensajes.append((
            "aviso",
            f"{int(a_perdida.sum())} plato(s) tienen precio menor al costo ({', '.join(nombres)}). "
            "Revisá que no estén invertidas las columnas de costo y precio.",
        ))
    if (out["units_sold"] == 0).all():
        mensajes.append(
            ("aviso", "Todas las unidades vendidas son 0: sin volumen no se puede calcular la popularidad.")
        )

    columnas = ["item_name", "category", "cost", "price", "units_sold"]
    return out[columnas].reset_index(drop=True), mensajes


# ---------------------------------------------------------------------------
# Plantilla para los alumnos
# ---------------------------------------------------------------------------
EJEMPLOS = [
    ("Milanesa napolitana", "Platos principales", 4.50, 16.00, 180),
    ("Bife de chorizo", "Platos principales", 11.00, 26.00, 60),
    ("Risotto de hongos", "Platos principales", 5.00, 18.00, 25),
    ("Ensalada César", "Entradas", 3.00, 9.50, 150),
    ("Provoleta", "Entradas", 2.80, 8.00, 95),
    ("Flan casero", "Postres", 1.20, 5.00, 40),
    ("Tiramisú", "Postres", 2.40, 7.50, 70),
]


def plantilla_bytes() -> bytes:
    """Genera el .xlsx que se les entrega a los alumnos para completar."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    ws = wb.active
    ws.title = "Menú"
    encabezados = ["Plato", "Categoría", "Costo", "Precio", "Unidades vendidas"]
    ws.append(encabezados)

    relleno = PatternFill("solid", fgColor="2563EB")
    for col, _ in enumerate(encabezados, start=1):
        celda = ws.cell(row=1, column=col)
        celda.font = Font(bold=True, color="FFFFFF", size=11)
        celda.fill = relleno
        celda.alignment = Alignment(horizontal="center", vertical="center")

    for fila in EJEMPLOS:
        ws.append(list(fila))

    for col, ancho in enumerate([30, 22, 12, 12, 20], start=1):
        ws.column_dimensions[get_column_letter(col)].width = ancho
    for fila in range(2, len(EJEMPLOS) + 2):
        ws.cell(row=fila, column=3).number_format = "#,##0.00"
        ws.cell(row=fila, column=4).number_format = "#,##0.00"
        ws.cell(row=fila, column=5).number_format = "#,##0"
    ws.freeze_panes = "A2"

    guia = wb.create_sheet("Instrucciones")
    texto = [
        ("Cómo completar esta planilla", True),
        ("", False),
        ("1. Borrá las filas de ejemplo de la hoja «Menú» y cargá tus propios platos.", False),
        ("2. Una fila por plato. No agregues filas de totales ni subtotales.", False),
        ("3. Categoría: agrupá los platos comparables entre sí (Entradas, Platos", False),
        ("   principales, Postres, Bebidas…). El análisis se hace DENTRO de cada", False),
        ("   categoría, porque no tiene sentido comparar una entrada con un plato", False),
        ("   principal: ni sus volúmenes ni sus márgenes son comparables.", False),
        ("4. Costo: lo que te cuesta producir una porción.", False),
        ("5. Precio: el precio de venta al público.", False),
        ("   El IVA se configura una sola vez en la app, en la ficha del restaurante.", False),
        ("   Cargá acá los valores tal como los tenés; la app netea el IVA sola.", False),
        ("6. Unidades vendidas: cantidad vendida EN TODO EL PERÍODO que vas a cargar.", False),
        ("", False),
        ("Sobre el período", True),
        ("El período debe abarcar semanas completas (7, 14, 21… días).", False),
        ("En gastronomía un sábado no vende como un martes: si el período tuviera", False),
        ("10 días incluiría dos sábados y un solo martes, y el mix quedaría sesgado.", False),
        ("", False),
        ("Podés cambiar los nombres de las columnas: la app reconoce las variantes", False),
        ("más comunes (Producto, Cantidad, Precio de venta, Food cost, etc.).", False),
    ]
    for i, (linea, negrita) in enumerate(texto, start=1):
        celda = guia.cell(row=i, column=1, value=linea)
        if negrita:
            celda.font = Font(bold=True, size=12, color="1F2937")
    guia.column_dimensions["A"].width = 95

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
